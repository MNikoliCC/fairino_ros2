#!/usr/bin/env python3
"""Build the Excel BOM for the robot rail base (formulas stay live in Excel) and,
optionally, refresh the BOM table in the README between the BOM markers.

    python gen_bom.py ../robot_rail_v3_BOM.xlsx [--readme ../README.md]
"""
import argparse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from bom_data import EUR_RSD, EUR_RSD_DATE, ITEMS, NOTES, VAT

ap = argparse.ArgumentParser()
ap.add_argument("xlsx")
ap.add_argument("--readme", help="README.md with <!-- BOM:START --> / <!-- BOM:END --> markers")
args = ap.parse_args()
OUT = args.xlsx

wb = Workbook()
ws = wb.active
ws.title = "BOM"

HDR = ["Rb.", "Šifra", "Naziv proizvoda", "Specifikacija / opis", "Dobavljač (Srbija)",
       "Link za nabavku", "Jed. mere", "Cena po jedinici", "Valuta", "Cena sadrži PDV?",
       "Jed. cena RSD bez PDV", "Količina", "Ukupno RSD bez PDV", "Ukupno RSD sa PDV",
       "Ukupno EUR bez PDV", "3D model (izvor)", "U 3D modelu", "Napomena"]
WIDTH = [5, 14, 34, 52, 18, 44, 8, 11, 7, 9, 13, 9, 15, 15, 13, 40, 10, 60]

thin = Side(style="thin", color="B0B7C3")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="1F3A5F")
sec_fill = PatternFill("solid", fgColor="DCE6F2")
tot_fill = PatternFill("solid", fgColor="FFF2CC")
wrap = Alignment(wrap_text=True, vertical="top")

ws["A1"] = "BOM – Robot linear rail (7. osa), v3"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = "Kurs EUR→RSD (NBS srednji):"
ws["E2"] = EUR_RSD
ws["F2"] = f"NBS srednji kurs na dan {EUR_RSD_DATE} – promenite ovde, sve se preračunava"
ws["A3"] = "PDV:"
ws["E3"] = VAT
ws["E3"].number_format = "0%"
for c in ("E2", "E3"):
    ws[c].fill = tot_fill
    ws[c].font = Font(bold=True)
KURS, PDV = "$E$2", "$E$3"

r0 = 5
for i, h in enumerate(HDR, 1):
    c = ws.cell(row=r0, column=i, value=h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = head_fill
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    c.border = border
    ws.column_dimensions[get_column_letter(i)].width = WIDTH[i - 1]
ws.row_dimensions[r0].height = 32

row = r0 + 1
n = 0
first_item_row = row
for it in ITEMS:
    if "section" in it:
        ws.cell(row=row, column=1, value=it["section"]).font = Font(bold=True)
        for col in range(1, len(HDR) + 1):
            ws.cell(row=row, column=col).fill = sec_fill
            ws.cell(row=row, column=col).border = border
        row += 1
        continue
    n += 1
    vals = [n, it["code"], it["name"], it["spec"], it["supplier"], it["url"], it["unit"],
            it["price"], it["cur"], "da" if it["vat_incl"] else "ne", None, it["qty"],
            None, None, None, it.get("cad", ""), it.get("in_model", "da"), it.get("note", "")]
    for col, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=col, value=v)
        c.alignment = wrap
        c.border = border
    # unit price in RSD without VAT
    ws.cell(row=row, column=11, value=(
        f'=IF(I{row}="EUR",H{row}*{KURS},H{row})/IF(J{row}="da",1+{PDV},1)'))
    ws.cell(row=row, column=13, value=f"=K{row}*L{row}")
    ws.cell(row=row, column=14, value=f"=M{row}*(1+{PDV})")
    ws.cell(row=row, column=15, value=f"=M{row}/{KURS}")
    for col in (8,):
        ws.cell(row=row, column=col).number_format = "#,##0.00"
    for col in (11, 13, 14):
        ws.cell(row=row, column=col).number_format = "#,##0.00"
    ws.cell(row=row, column=15).number_format = "#,##0.00 €"
    if it["url"].startswith("http"):
        ws.cell(row=row, column=6).hyperlink = it["url"]
        ws.cell(row=row, column=6).font = Font(color="1F5FBF", underline="single")
    cad = it.get("cad", "")
    if cad.startswith("http"):
        ws.cell(row=row, column=16).hyperlink = cad
        ws.cell(row=row, column=16).font = Font(color="1F5FBF", underline="single")
    row += 1
last_item_row = row - 1

row += 1
tot_row = row
for label, col, fmt in (("UKUPNO bez PDV (RSD)", 13, "#,##0.00"),
                        ("UKUPNO sa PDV (RSD)", 14, "#,##0.00"),
                        ("UKUPNO bez PDV (EUR)", 15, "#,##0.00 €")):
    ws.cell(row=row, column=3, value=label).font = Font(bold=True, size=12)
    c = ws.cell(row=row, column=col,
                value=f"=SUM({get_column_letter(col)}{first_item_row}:{get_column_letter(col)}{last_item_row})")
    c.font = Font(bold=True, size=12)
    c.number_format = fmt
    for cc in range(1, len(HDR) + 1):
        ws.cell(row=row, column=cc).fill = tot_fill
        ws.cell(row=row, column=cc).border = border
    row += 1

row += 1
ws.cell(row=row, column=1, value="Napomene").font = Font(bold=True, size=12)
row += 1
for t in NOTES:
    ws.cell(row=row, column=1, value="•")
    ws.cell(row=row, column=2, value=t).alignment = Alignment(wrap_text=False, vertical="top")
    row += 1

# summary block at the top
for k, (label, col, fmt) in enumerate((("UKUPNO bez PDV (RSD):", 13, "#,##0.00"),
                                       ("UKUPNO sa PDV (RSD):", 14, "#,##0.00"),
                                       ("UKUPNO bez PDV (EUR):", 15, "#,##0.00 €"))):
    ws.cell(row=2 + k, column=8, value=label).font = Font(bold=True)
    c = ws.cell(row=2 + k, column=11, value=f"={get_column_letter(col)}{tot_row + k}")
    c.font = Font(bold=True)
    c.number_format = fmt
    c.fill = tot_fill

ws.freeze_panes = ws.cell(row=r0 + 1, column=4)
ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(HDR))}{last_item_row}"
ws.sheet_view.zoomScale = 90
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT, "items", n)


def unit_rsd_net(it):
    p = it["price"]
    if isinstance(p, str):                      # simple constant formula, e.g. "=0.1*2.8*10.19"
        p = eval(p.lstrip("="), {"__builtins__": {}})
    p = p * EUR_RSD if it["cur"] == "EUR" else p
    return p / (1 + VAT) if it["vat_incl"] else p


def readme_table():
    def money(v):
        return f"{v:,.2f}".replace(",", " ").replace(".", ",")
    rows = ["| # | Šifra | Naziv | Dobavljač | Kol. | Jed. | Cena / jed. | Ukupno RSD bez PDV | Ukupno RSD sa PDV |",
            "|---|---|---|---|---:|---|---:|---:|---:|"]
    n, total = 0, 0.0
    for it in ITEMS:
        if "section" in it:
            rows.append(f"| | | **{it['section']}** | | | | | | |")
            continue
        n += 1
        line = unit_rsd_net(it) * it["qty"]
        total += line
        price = it["price"]
        price = unit_rsd_net(it) if isinstance(price, str) else price
        vat = "sa PDV" if it["vat_incl"] else "bez PDV"
        name = f"[{it['name']}]({it['url']})" if it["url"].startswith("http") else it["name"]
        qty = f"{it['qty']:g}".replace(".", ",")
        rows.append(f"| {n} | {it['code']} | {name} | {it['supplier']} | {qty} | {it['unit']} | "
                    f"{money(price)} {it['cur'] if not isinstance(it['price'], str) else 'RSD'} {vat} | {money(line)} | "
                    f"{money(line * (1 + VAT))} |")
    rows.append(f"| | | **UKUPNO (RSD)** | | | | | **{money(total)}** | **{money(total * (1 + VAT))}** |")
    rows.append(f"| | | **UKUPNO (EUR, kurs {EUR_RSD})** | | | | | **{money(total / EUR_RSD)} €** | "
                f"**{money(total * (1 + VAT) / EUR_RSD)} €** |")
    summary = ("| | bez PDV | sa PDV ({:.0%}) |\n|---|---:|---:|\n"
               "| **Ukupno RSD** | **{}** | **{}** |\n| **Ukupno EUR** (kurs {}) | {} € | {} € |\n\n"
               "CNC obrada ploče kolica nije uračunata (po ponudi).\n").format(
        VAT, money(total), money(total * (1 + VAT)), EUR_RSD, money(total / EUR_RSD),
        money(total * (1 + VAT) / EUR_RSD))
    return summary + "\n" + "\n".join(rows)


if args.readme:
    text = open(args.readme, encoding="utf-8").read()
    a, z = "<!-- BOM:START -->", "<!-- BOM:END -->"
    i, j = text.index(a) + len(a), text.index(z)
    open(args.readme, "w", encoding="utf-8").write(text[:i] + "\n" + readme_table() + "\n" + text[j:])
    print("README BOM table updated:", args.readme)
